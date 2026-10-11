import base64
import asyncio
import httpx
import io
import json
import os
import re
import secrets
import threading
import time
from datetime import datetime,timezone,timedelta
from pathlib import Path
from fastapi import FastAPI,Request,HTTPException,UploadFile,File
from fastapi.responses import FileResponse,Response
from fastapi.staticfiles import StaticFiles
from core import *
from store import Store,password_ok,totp
from profile import load_profile
from ai import assist,readiness
from templates import catalog,load_template

HERE=Path(__file__).resolve().parent
store=Store(os.environ.get('ESCRIBANIA_DATA',str(HERE/'runtime')))
GLOBAL_PROFILE=load_profile(store.directory)
PRODUCTION=os.environ.get('ESCRIBANIA_MODE','local')=='production'
ORIGIN=os.environ.get('ESCRIBANIA_ORIGIN','http://127.0.0.1:8001').rstrip('/')
if PRODUCTION and not ORIGIN.startswith('https://'): raise RuntimeError('Producción requiere HTTPS y ESCRIBANIA_ORIGIN')
app=FastAPI(title='EscribanIA',docs_url=None,redoc_url=None)
sessions={}; failures={}; used_totp={}; lock=threading.RLock()
request_lock=asyncio.Lock()
SOURCES=[{'title':'Ley 18.331 — protección de datos personales','url':'https://www.impo.com.uy/bases/leyes/18331-2008'}, {'title':'Decreto 414/009 — reglamentación','url':'https://www.impo.com.uy/bases/decretos/414-2009'}, {'title':'Ley 18.600 — documento y firma electrónicos','url':'https://www.impo.com.uy/bases/leyes/18600-2009'}, {'title':'Acordada 7.533 — Reglamento Notarial; verificar modificaciones','url':'https://www.impo.com.uy/bases/acordadas-scj-originales/7533-2004'}, {'title':'URCDP — orientación oficial','url':'https://www.gub.uy/unidad-reguladora-control-datos-personales/'}]

@app.middleware('http')
async def security(request,call_next):
    if request.method not in ('GET','HEAD','OPTIONS') and request.url.path!='/api/login':
        if request.headers.get('origin')!=ORIGIN: return Response('Origen no permitido',403)
        session=sessions.get(request.cookies.get('escribania',''))
        if not session or not secrets.compare_digest(request.headers.get('x-csrf-token',''),session['csrf']): return Response('CSRF o sesión inválidos',403)
    if request.url.path=='/api/login' and request.headers.get('origin')!=ORIGIN: return Response('Origen no permitido',403)
    async with request_lock:
        response=await call_next(request)
    response.headers['Cache-Control']='no-store'
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Referrer-Policy']='no-referrer'
    response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; object-src 'none'; base-uri 'none'; frame-ancestors 'none'; connect-src 'self'"
    if PRODUCTION: response.headers['Strict-Transport-Security']='max-age=31536000'
    return response

def user(request):
    session=sessions.get(request.cookies.get('escribania',''))
    if not session or session['expires']<time.time(): raise HTTPException(401,'Iniciá sesión con contraseña y TOTP')
    return session['user']

def case_for(cid,request):
    try: return store.get(cid,user(request))
    except PermissionError: raise HTTPException(404,'Expediente no disponible')
    except ValueError: raise HTTPException(409,'Integridad de auditoría inválida; trabajo bloqueado')

def editable(case):
    if case['approved']: raise HTTPException(409,'Reabrí el expediente con motivo antes de modificarlo')

def working(case):
    chosen={c['segment']:c['after'] for c in case['changes'] if c['status']=='accepted'}
    return [dict(s,text=chosen.get(s['id'],s['text'])) for s in case['segments']]

def invalidate(case):
    case['corroboration']=None; case['analysis']=None
    case.pop('reviewed_document',None); case.pop('reviewed_hash',None)

def blockers(case):
    result=identity_issues(analyzed_text(working(case)),case['parties'])
    if re.search(r'\[[^\]\n]+\]',analyzed_text(working(case)),re.I): result.append('Campos entre corchetes pendientes de completar')
    result+=['Cambio pendiente: '+c['id'] for c in case['changes'] if c['status']=='pending']
    result+=['Observación sin decisión: '+o['id'] for o in case.get('observations',[]) if not o.get('decision')]
    if not case['corroboration']: result.append('Falta cotejo integral de todas las menciones y asociaciones de identidad con evidencia')
    if not case['analysis']: result.append('Falta análisis actualizado')
    if case['warnings']: result.append('Elementos Word no procesados o complejos: requieren tratamiento fuera de esta versión')
    if not GLOBAL_PROFILE: result.append('Perfil normativo único: dimensiones y márgenes pendientes de plantilla técnica validada')
    elif case['template']!=GLOBAL_PROFILE: result.append('Perfil normativo actualizado: registrá análisis y corroboración de nuevo')
    return result

def view(case):
    return {k:v for k,v in case.items() if k not in ('original','final_document','reviewed_document')} | {'working':working(case),'blockers':blockers(case),'audit_status':store.verify()}

def reason(data):
    value=str(data.get('reason','')).strip(); evidence=str(data.get('evidence','')).strip()
    if len(value)<12 or len(evidence)<8: raise HTTPException(422,'Indicá fundamento (12 caracteres) y evidencia consultada (8 caracteres)')
    return value,evidence

def save(case,request,action,details): store.save(case,user(request),action,details)

@app.get('/api/status')
def status(): return {'version':VERSION,'mode':'production' if PRODUCTION else 'local','ai':'deshabilitada hasta configurar y validar proveedor','sources':SOURCES,'source_status':'Consulta oficial documentada en docs/fuentes.md; verificar consolidación y modificaciones antes del uso','global_profile':GLOBAL_PROFILE,'publication':'No publicada por esta implementación'}

@app.post('/api/login')
async def login(request:Request):
    data=await request.json(); name=str(data.get('user','')); ip=request.client.host if request.client else 'unknown'; now=time.time()
    # Both account and IP rate limits; no credential/document logging.
    for key in ('user:'+name,'ip:'+ip):
        failures[key]=[t for t in failures.get(key,[]) if now-t<900]
        if len(failures[key])>=8: raise HTTPException(429,'Demasiados intentos; esperá 15 minutos')
    with store.connect() as db: row=db.execute('SELECT * FROM users WHERE name=?',(name,)).fetchone()
    valid=False; matched=None
    if row:
        secret=store.fernet.decrypt(row['totp']).decode()
        matched=next((int(now)//30+d for d in (-1,0,1) if secrets.compare_digest(totp(secret,now+30*d),str(data.get('otp','')))),None)
        valid=password_ok(str(data.get('password','')),row['password']) and matched is not None and matched>used_totp.get(name,-1)
    if not valid:
        for key in ('user:'+name,'ip:'+ip): failures[key].append(now)
        raise HTTPException(401,'Credenciales o código incorrectos')
    used_totp[name]=matched; token=secrets.token_urlsafe(32); csrf=secrets.token_urlsafe(32)
    sessions[token]={'user':name,'expires':now+1800,'csrf':csrf}
    response=Response(json.dumps({'user':name,'csrf':csrf}),media_type='application/json')
    response.set_cookie('escribania',token,secure=PRODUCTION,httponly=True,samesite='strict',max_age=1800)
    return response

@app.get('/api/me')
def me(request:Request): return {'user':user(request),'csrf':sessions[request.cookies['escribania']]['csrf']}
@app.post('/api/logout')
def logout(request:Request):
    sessions.pop(request.cookies.get('escribania',''),None); response=Response('{}',media_type='application/json'); response.delete_cookie('escribania'); return response
@app.get('/api/cases')
def list_cases(request:Request):
    name=user(request)
    with store.connect() as db:
        return [{'id':r['id'],'title':store.decrypt(r['payload'])['title']} for r in db.execute('SELECT * FROM cases WHERE owner=? OR EXISTS(SELECT 1 FROM acl WHERE case_id=cases.id AND user=?)',(name,name))]

@app.post('/api/cases')
async def create_case(request:Request,file:UploadFile=File(...)):
    name=user(request)
    if not file.filename.lower().endswith('.docx'): raise HTTPException(422,'Importá un archivo .docx')
    blob=await file.read(10*1024*1024+1)
    return import_document(blob,file.filename,name)

def import_document(blob,filename,name,template_origin=None):
    try: segments,warnings=read_docx(blob)
    except ValueError as e: raise HTTPException(422,str(e))
    cid=secrets.token_hex(12)
    case={'id':cid,'title':filename,'original':base64.b64encode(blob).decode(),'original_hash':digest(blob),'text_hash':digest(analyzed_text(segments)), 'segments':segments,'warnings':warnings,'parties':[],'changes':[],'observations':[],'analysis':None,'corroboration':None,'template':None,'approved':False,'final_hash':None,'retention_until':(datetime.now(timezone.utc)+timedelta(days=30)).isoformat(),'versions':{'app':VERSION,'rules':RULES_VERSION,'rules_sha256':digest((HERE/'core.py').read_bytes()),'prompt':PROMPT_VERSION,'model':None},'sources':[]}
    if template_origin: case['template_origin']=template_origin
    store.save(case,name,'template_case_created' if template_origin else 'import',{'original_sha256':case['original_hash'],'text_sha256':case['text_hash'],'warnings':warnings,'template_origin':template_origin},create=True)
    return view(case)

@app.get('/api/templates')
def templates(request:Request):
    user(request)
    try: return catalog()
    except ValueError as e: raise HTTPException(409,str(e))

def template_data(template_id):
    try: return load_template(template_id)
    except KeyError: raise HTTPException(404,'Plantilla no disponible')
    except ValueError as e: raise HTTPException(409,str(e))

@app.get('/api/templates/{template_id}/download')
def template_download(template_id:str,request:Request):
    user(request); manifest,blob=template_data(template_id)
    return Response(blob,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="EscribanIA-{template_id}-borrador.docx"'})

@app.post('/api/templates/{template_id}/cases')
async def template_case(template_id:str,request:Request):
    name=user(request); data=await request.json()
    if data.get('review_required_acknowledged') is not True: raise HTTPException(422,'Confirmá que revisaste las adaptaciones y que usarás el documento como borrador pendiente de revisión profesional')
    manifest,blob=template_data(template_id)
    origin={k:manifest[k] for k in ('id','version','source_sha256','converted_sha256','template_sha256','review_required','normative_format_verified')}
    origin['adaptations_sha256']=digest(json.dumps(manifest['changes'],ensure_ascii=False,sort_keys=True))
    if 'sources' in manifest: origin['sources']=manifest['sources']
    if 'source_hash_scope' in manifest: origin['source_hash_scope']=manifest['source_hash_scope']
    if 'legal_review' in manifest: origin['legal_review']=manifest['legal_review']
    return import_document(blob,manifest['title']+' — borrador.docx',name,origin)

@app.get('/api/cases/{cid}')
def get_case(cid:str,request:Request): return view(case_for(cid,request))

@app.post('/api/draft')
async def draft(request:Request):
    user(request); data=await request.json()
    title=str(data.get('title','Borrador para revisión'))[:150]
    if not title.strip(): raise HTTPException(422,'Objeto requerido')
    from docx import Document
    doc=Document(); doc.add_heading('BORRADOR — NO APROBADO',0); doc.add_paragraph(title)
    for text in ['Comparecientes: [PENDIENTE: fichas cotejadas e identidades completas].','Antecedentes y objeto: [PENDIENTE: documentos y hechos corroborados].','Normativa aplicable: [PENDIENTE: fuente oficial vigente y artículos verificados].','Cláusulas: [PENDIENTE: redacción profesional según el acto concreto].','Viabilidad: INDETERMINADA. Revisar capacidad, representación, legitimación, antecedentes, restricciones y requisitos del acto. No sustituye el criterio profesional.']:
        doc.add_paragraph(text)
    out=io.BytesIO(); doc.save(out)
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':'attachment; filename="EscribanIA-borrador-base.docx"'})

@app.post('/api/cases/{cid}/source')
async def source(cid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    url=str(data.get('url',''))
    if not url.startswith(('https://www.impo.com.uy/','https://www.gub.uy/','https://www.poderjudicial.gub.uy/')): raise HTTPException(422,'Registrá una fuente oficial')
    entry={'url':url,'title':str(data.get('title',''))[:250],'article':str(data.get('article',''))[:250],'excerpt':str(data.get('excerpt',''))[:10000],'consulted_at':datetime.now(timezone.utc).isoformat(),'reviewer':user(request),'reason':why,'evidence':evidence,'verification':'Registrada por revisor; no descargada automáticamente por el sistema'}
    if not entry['title'] or not entry['excerpt']: raise HTTPException(422,'Título y fragmento necesarios')
    case['sources'].append(entry); invalidate(case); save(case,request,'source_recorded',entry); return view(case)

@app.post('/api/cases/{cid}/parties')
async def add_party(cid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    name=upper_name(str(data.get('name',''))); ci=str(data.get('ci','')).strip()
    if len(name.split())<2 or len(name)>200 or not re.fullmatch(r"[A-ZÁÉÍÓÚÜÑÀÈÌÒÙÂÊÎÔÛÇ '’-]+",name): raise HTTPException(422,'Ingresá nombres y apellidos completos, sin números ni referencias')
    if not valid_ci(ci): raise HTTPException(422,'Formato o dígito verificador inválido; una validación correcta no acredita identidad')
    if any(p['name']==name or ci_digits(p['ci'])==ci_digits(ci) for p in case['parties']): raise HTTPException(409,'Ficha contradictoria o duplicada: revisá los datos')
    party={'id':secrets.token_hex(8),'name':name,'ci':ci,'verified_by':user(request),'evidence':evidence,'reason':why}
    case['parties'].append(party); invalidate(case); save(case,request,'identity_verified',party)
    return view(case)

@app.post('/api/cases/{cid}/changes')
async def propose(cid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    segment=next((s for s in case['segments'] if s['id']==data.get('segment')),None)
    if not segment or segment['deleted']: raise HTTPException(422,'Segmento no editable')
    after=str(data.get('after',''))
    if len(after)>20000 or after==segment['text']: raise HTTPException(422,'Cambio vacío o excesivo')
    if any(c['segment']==segment['id'] and c['status']!='undone' for c in case['changes']): raise HTTPException(409,'Deshacé primero el cambio existente sobre este segmento')
    # The identity set may not be added, removed or altered. Case-only names are permitted solely to restore canonical uppercase.
    before=analyzed_text(working(case))
    after_full=analyzed_text([dict(s,text=after if s['id']==segment['id'] else s['text']) for s in working(case)])
    before_ci=[m[0] for m in CI.finditer(before)]; after_ci=[m[0] for m in CI.finditer(after_full)]
    if before_ci!=after_ci: raise HTTPException(409,'Las cédulas no pueden modificarse en propuestas')
    for party in case['parties']:
        pattern=r'(?<!\w)'+re.escape(party['name'])+r'(?!\w)'
        old=list(re.finditer(pattern,before,re.I)); new=list(re.finditer(pattern,after_full,re.I))
        if len(old)!=len(new) or any(m[0]!=party['name'] for m in new): raise HTTPException(409,'Identidad protegida: solo se admite la forma cotejada en MAYÚSCULAS')
    change={'id':secrets.token_hex(8),'segment':segment['id'],'part':segment['part'],'index':segment['index'],'before':segment['text'],'after':after,'reason':why,'evidence':evidence,'status':'pending','author':user(request)}
    case['changes'].append(change); invalidate(case); save(case,request,'change_proposed',change)
    return view(case)

@app.post('/api/cases/{cid}/changes/{change_id}')
async def decide_change(cid:str,change_id:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    change=next((c for c in case['changes'] if c['id']==change_id),None); status=data.get('status')
    if not change or status not in ('accepted','rejected','undone'): raise HTTPException(422,'Decisión inválida')
    if change['status']=='undone' or (change['status']!='pending' and status!='undone'): raise HTTPException(409,'Transición inválida')
    change['status']=status; change['reviewer']=user(request); invalidate(case)
    save(case,request,'change_'+status,{'change':change_id,'before':change['before'],'after':change['after'],'reason':why,'evidence':evidence})
    return view(case)

@app.post('/api/cases/{cid}/identity-fix')
async def identity_fix(cid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    party=next((p for p in case['parties'] if p['id']==data.get('party')),None)
    segment=next((s for s in case['segments'] if s['id']==data.get('segment') and not s['deleted']),None)
    field=data.get('field'); old=str(data.get('old',''))
    if not party or not segment or field not in ('name','ci') or len(old)<4 or segment['text'].count(old)!=1: raise HTTPException(422,'Elegí ficha cotejada, campo y una coincidencia exacta del original')
    if old==party[field]: raise HTTPException(422,'El dato ya coincide')
    if any(c['segment']==segment['id'] and c['status']!='undone' for c in case['changes']): raise HTTPException(409,'Deshacé primero la propuesta existente')
    for p in case['parties']:
        if field=='name' and old in p['name']: raise HTTPException(409,'El dato pertenece a una identidad protegida; no se reasigna')
        if field=='ci':
            try:
                if ci_digits(old)==ci_digits(p['ci']): raise HTTPException(409,'Cédula ya cotejada; no se reasigna')
            except ValueError: pass
    after=segment['text'].replace(old,party[field],1)
    change={'id':secrets.token_hex(8),'segment':segment['id'],'part':segment['part'],'index':segment['index'],'before':segment['text'],'after':after,'reason':why,'evidence':evidence,'status':'pending','author':user(request),'type':'deterministic_identity_correction','party':party['id'],'field':field}
    case['changes'].append(change);invalidate(case);save(case,request,'identity_correction_proposed',change);return view(case)

@app.post('/api/cases/{cid}/analyze')
async def analyze(cid:str,request:Request):
    case=case_for(cid,request); editable(case)
    case['template']=GLOBAL_PROFILE
    text=analyzed_text(working(case)); issues=identity_issues(text,case['parties'])
    observations=[]
    for issue in issues:
        observations.append({'id':secrets.token_hex(8),'kind':'hecho','finding':issue,'fragment':text[:300] if ':' not in issue else issue.split(':',1)[1].strip(),'basis':'Comparación determinista con fichas cotejadas y regla de dígito verificador; no acredita identidad.', 'alternatives':'Error de transcripción, compareciente omitido o documentación insuficiente.','uncertainty':'No permite decidir autenticidad ni identidad real.','reconsider':'Documento de identidad cotejado y revisión íntegra de todas las menciones.','decision':None})
    for pattern,label in [(r'\[[^\]\n]+\]','Campo entre corchetes pendiente'),(r'\b([^\W\d_]+)\s+\1\b','Posible palabra repetida')]:
        for match in list(re.finditer(pattern,text,re.I))[:20]:
            observations.append({'id':secrets.token_hex(8),'kind':'hecho','finding':label,'fragment':match[0],'basis':'Control mecánico portado de lib/review.ts del prototipo v0.3; coincidencia literal.','alternatives':'Marcador intencional o repetición admisible; revisar contexto.','uncertainty':'La coincidencia no demuestra un error jurídico.','reconsider':'Texto completo, finalidad del campo y documentación del acto.','decision':None})
    observations.append({'id':secrets.token_hex(8),'kind':'indeterminable','finding':'Viabilidad jurídica y vigencia normativa no determinadas','fragment':'Documento completo','basis':'No existe consulta a modelo ni corpus oficial actualizado verificado en esta instancia.','alternatives':'Consulta fundada de fuentes oficiales y análisis de la escribana.','uncertainty':'No se emite conclusión de viabilidad ni se completan datos faltantes.','reconsider':'Identificar acto, jurisdicción, antecedentes, fuente vigente y sus artículos aplicables.','decision':None})
    case['observations']=observations; case['analysis']={'sha256':digest(text),'at':datetime.now(timezone.utc).isoformat(),'engine':'reglas deterministas; IA deshabilitada','versions':case['versions'],'sources':case['sources']}
    try:
        reviewed=export_docx(base64.b64decode(case['original']),case['changes'],user(request),GLOBAL_PROFILE)
        case['reviewed_document']=base64.b64encode(reviewed).decode();case['reviewed_hash']=digest(reviewed)
    except ValueError:
        case.pop('reviewed_document',None);case.pop('reviewed_hash',None)
    case['corroboration']=None; save(case,request,'analysis',{'analysis':case['analysis'],'observations':observations})
    return view(case)

@app.post('/api/cases/{cid}/observations/{oid}')
async def observation_decision(cid:str,oid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    obs=next((o for o in case['observations'] if o['id']==oid),None)
    if not obs or data.get('decision') not in ('accepted','rejected','challenged'): raise HTTPException(422,'Decisión inválida')
    obs['decision']=data['decision']; obs['reviewer']=user(request); obs['reason']=why; obs['evidence']=evidence
    if data['decision']=='challenged': obs['decision']=None
    save(case,request,'observation_'+data['decision'],{'observation':oid,'reason':why,'evidence':evidence})
    return view(case)

@app.post('/api/cases/{cid}/corroborate')
async def corroborate(cid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    if not case['analysis']: raise HTTPException(409,'Analizá la versión actual primero')
    if data.get('all_mentions_checked') is not True: raise HTTPException(422,'Confirmá cotejo íntegro, incluidas asociaciones nombre-cédula y variantes')
    if GLOBAL_PROFILE and (data.get('print_checked') is not True or data.get('reviewed_sha256')!=case.get('reviewed_hash')): raise HTTPException(422,'Revisá el Word de esta huella e impresión: 25 líneas/carilla, 55 caracteres/línea y separación mínima 8 mm, anverso y reverso')
    case['corroboration']={'reviewer':user(request),'reason':why,'evidence':evidence,'text_sha256':digest(analyzed_text(working(case))),'print_checked':data.get('print_checked') is True,'reviewed_sha256':data.get('reviewed_sha256')}
    save(case,request,'corroborated',case['corroboration']); return view(case)

@app.post('/api/cases/{cid}/pseudonymize')
async def pseudonym_preview(cid:str,request:Request):
    case=case_for(cid,request); data=await request.json()
    try: text,mapping,residuals=pseudonymize(analyzed_text(working(case)),case['parties'],data.get('extra_terms',[]))
    except ValueError as e: raise HTTPException(422,str(e))
    save(case,request,'local_pseudonym_preview',{'residual_categories':residuals,'outbound':False})
    return {'text':text,'sha256':digest(text),'residuals':residuals,'outbound':False,'notice':'Seudonimización reversible. La revisión de identificadores indirectos sigue siendo necesaria. Esta previsualización no envía contenido fuera del servidor.'}

@app.post('/api/cases/{cid}/assist')
async def model_assist(cid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json()
    if readiness(store.directory): raise HTTPException(503,readiness(store.directory))
    text=analyzed_text(working(case))
    if identity_issues(text,case['parties']): raise HTTPException(409,'Resolvé discrepancias de identidad antes de enviar')
    try:
        from starlette.concurrency import run_in_threadpool
        result=await run_in_threadpool(assist,store.directory,text,case['parties'],data.get('extra_terms',[]),data.get('module'),data.get('preview_hash'),data.get('reviewed') is True)
    except (ValueError,httpx.HTTPError) as e: raise HTTPException(422,str(e) if isinstance(e,ValueError) else 'Error de conexión al proveedor')
    for observation in result['observations']:
        observation.update(id=secrets.token_hex(8),decision=None)
        case['observations'].append(observation)
    case['corroboration']=None; case['analysis']={'sha256':digest(text),'engine':'IA + reglas; revisión pendiente','metadata':result['metadata']}
    case['versions']['model']=result['metadata']['model'];case['versions']['prompt']=result['metadata']['prompt_version']
    save(case,request,'model_assistance',{'metadata':result['metadata'],'observations':result['observations'],'draft_sha256':digest(result.get('draft',''))})
    return {'case':view(case),'draft':result.get('draft',''),'metadata':result['metadata']}

@app.post('/api/cases/{cid}/approve')
async def approve(cid:str,request:Request):
    case=case_for(cid,request); editable(case); data=await request.json(); why,evidence=reason(data)
    problems=blockers(case)
    if problems: raise HTTPException(409,problems)
    try:
        final=base64.b64decode(case['reviewed_document'])
        if digest(final)!=case['reviewed_hash']: raise ValueError('La versión revisada fue alterada')
        if analyzed_text(read_docx(final)[0])!=analyzed_text(working(case)): raise ValueError('El documento exportado no coincide con el texto revisado')
    except ValueError as e: raise HTTPException(409,str(e))
    case['approved']=True; case['final_hash']=digest(final); case['final_document']=base64.b64encode(final).decode()
    save(case,request,'approved',{'final_sha256':case['final_hash'],'text_sha256':digest(analyzed_text(working(case))),'reason':why,'evidence':evidence,'template':case['template']})
    return view(case)

@app.post('/api/cases/{cid}/reopen')
async def reopen(cid:str,request:Request):
    case=case_for(cid,request); data=await request.json(); why,evidence=reason(data)
    if not case['approved']: raise HTTPException(409,'No está aprobado')
    case['approved']=False; case['final_hash']=None; case.pop('final_document',None); invalidate(case)
    save(case,request,'reopened',{'reason':why,'evidence':evidence}); return view(case)

@app.get('/api/cases/{cid}/download/{kind}')
def download(cid:str,kind:str,request:Request):
    case=case_for(cid,request)
    if kind=='original': blob=base64.b64decode(case['original'])
    elif kind=='final':
        if not case['approved']: raise HTTPException(409,'Aprobación final pendiente')
        blob=base64.b64decode(case['final_document'])
        if digest(blob)!=case['final_hash']: raise HTTPException(409,'Integridad final inválida')
    elif kind=='draft':
        try: blob=base64.b64decode(case['reviewed_document']) if case.get('reviewed_document') else export_docx(base64.b64decode(case['original']),case['changes'],user(request))
        except ValueError as e: raise HTTPException(409,str(e))
    else: raise HTTPException(404)
    save(case,request,'download_'+kind,{'sha256':digest(blob)})
    return Response(blob,media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':f'attachment; filename="EscribanIA-{kind}.docx"'})

@app.get('/api/cases/{cid}/audit')
def audit(cid:str,request:Request): return {'verification':store.verify(),'events':store.history(cid,user(request))}

@app.post('/api/cases/{cid}/challenge')
async def challenge(cid:str,request:Request):
    case=case_for(cid,request); data=await request.json(); why,evidence=reason(data)
    case['approved']=False; case['final_hash']=None; case.pop('final_document',None); invalidate(case)
    save(case,request,'system_challenge',{'reason':why,'evidence':evidence}); return view(case)

@app.post('/api/cases/{cid}/share')
async def share(cid:str,request:Request):
    case=case_for(cid,request); data=await request.json(); why,evidence=reason(data)
    with store.connect() as db:
        if db.execute('SELECT owner FROM cases WHERE id=?',(cid,)).fetchone()['owner']!=user(request): raise HTTPException(403,'Solo quien creó el expediente gestiona permisos')
        if not db.execute('SELECT name FROM users WHERE name=?',(data.get('user'),)).fetchone(): raise HTTPException(422,'Usuario no registrado')
        db.execute('INSERT OR IGNORE INTO acl VALUES(?,?)',(cid,data['user']))
        store.event(db,cid,user(request),'permission_granted',{'user':data['user'],'reason':why,'evidence':evidence})
    return {'ok':True}

@app.post('/api/cases/{cid}/delete')
async def delete(cid:str,request:Request):
    case=case_for(cid,request); data=await request.json(); why,evidence=reason(data)
    with store.connect() as db:
        if db.execute('SELECT owner FROM cases WHERE id=?',(cid,)).fetchone()['owner']!=user(request): raise HTTPException(403,'Solo titular del expediente')
        store.event(db,cid,user(request),'deleted',{'reason':why,'evidence':evidence,'policy':'Borrado lógico del expediente; auditoría cifrada retenida. Backups y borrado criptográfico requieren operación externa.'})
        db.execute('DELETE FROM acl WHERE case_id=?',(cid,)); db.execute('DELETE FROM cases WHERE id=?',(cid,))
    return {'ok':True,'notice':'Expediente eliminado de la aplicación. No equivale a borrado físico de SQLite, backups o auditoría.'}

@app.get('/manual.pdf')
def manual(): return FileResponse(HERE/'docs/manual.pdf',media_type='application/pdf',filename='EscribanIA-manual-0.4.pdf')
app.mount('/static',StaticFiles(directory=HERE/'web'),name='static')
@app.get('/')
def index(): return FileResponse(HERE/'web/index.html')
