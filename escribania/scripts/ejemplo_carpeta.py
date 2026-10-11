"""Genera exclusivamente un ejemplo didáctico ficticio, sin documentos de clientes."""
from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from core import upper_name,valid_ci
ROOT=Path(__file__).resolve().parents[1]
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
# Formato de lectura A4: no representa medidas normativas de papel notarial.
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2)
normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(8)
s.header.paragraphs[0].text='ESCRIBANIA · EJEMPLO FICTICIO · SIN VALIDEZ PARA PRESENTACIÓN'
s.footer.paragraphs[0].text='Formato A4 de lectura. Márgenes ilustrativos; no acreditados como papel notarial.'
v=upper_name('María Muñoz Pérez');c=upper_name('José Núñez García');e=upper_name('Ana Ejemplo Notarial')
def ci(base):
 digit=(10-sum(int(a)*b for a,b in zip(base,[2,9,8,7,6,3,4]))%10)%10
 value=f'{base[0]}.{base[1:4]}.{base[4:]}-{digit}';assert valid_ci(value);return value
vci=ci('1234567');cci=ci('2345678')
def p(text):d.add_paragraph(text)
def title(n,text):
 if n:d.add_page_break()
 d.add_heading(text,0)
 d.add_paragraph('SIMULACIÓN DIDÁCTICA — Todos los datos, documentos, actuaciones y antecedentes que siguen son ficticios.').runs[0].font.color.rgb=RGBColor.from_string('9B3327')
def table(rows):
 t=d.add_table(rows=0,cols=2);t.style='Light Shading Accent 1'
 for k,val in rows:
  r=t.add_row();r.cells[0].text=k;r.cells[1].text=val
vehicle='vehículo automotor ficticio, marca MARCA EJEMPLO, modelo MODELO DEMO, año 2020, tipo sedán, combustible nafta, padrón DEMO-001 de Montevideo, matrícula DEMO-001, motor MOTOR-DEMO-001 y chasis CHASIS-DEMO-001'
title(0,'Carpeta de compraventa de automotor — ejemplo')
p('Versión de ejemplo · 11 de octubre de 2026. Permite ver la relación entre el contrato privado, su certificación, la protocolización y el testimonio que se presenta con carátula y minuta. No contiene firmas, sellos, signos ni evidencia de una actuación real.')
p('Las cédulas utilizadas son valores sintéticos con formato y dígito verificador comprobados. No acreditan identidad ni se afirma que estén asignadas a estos personajes. Ningún antecedente, póliza ni inscripción del ejemplo existe.')
table([('Parte vendedora',v+' · '+vci),('Parte compradora',c+' · '+cci),('Profesional del ejemplo',e+' · personaje ficticio, sin habilitación acreditada'),('Vehículo','DEMO-001 · todos sus identificadores son ficticios'),('Precio del supuesto','USD 7.800 — siete mil ochocientos dólares estadounidenses'),('Representación','Comparecencia personal supuesta; no se incluye poder')])
p('Índice: 1. Carátula ilustrativa. 2. Minuta ilustrativa. 3. Documento privado. 4. Certificación de firmas. 5. Acta de protocolización. 6. Estructura del testimonio notarial.')
p('Para una operación real, sustituir los datos mediante fichas cotejadas, corroborar antecedentes y recaudos, utilizar el formulario oficial de minuta y revisar el contenido y formato por la escribana. La generación de este archivo no cumple esas actuaciones.')
title(1,'1. Carátula envolvente — esquema')
p('Esquema de organización de la carpeta. No reproduce ni sustituye la carátula oficial ni acredita que se hayan completado todos los datos del artículo 2 del Decreto 58/995.')
table([('Registro de destino supuesto','Registro de la Propiedad, Sección Mobiliaria, Montevideo; primera inscripción allí supuesta'),('Acto','Compraventa de automotor'),('Documento portante','Testimonio notarial de protocolización del contrato privado y certificación de firmas'),('Otorgantes',v+' / '+c),('Vehículo','Padrón y matrícula DEMO-001, Montevideo'),('Escribana del supuesto',e),('Fecha del contrato','11 de octubre de 2026'),('Presentante y contacto','Persona presentante del ejemplo — contacto pendiente de cotejo'),('Número/fecha de presentación DGR','Sin presentación: no existe número de ingreso')])
p('Contenido para ventanilla: carátula, minuta oficial completa y suscrita, testimonio de protocolización expedido en forma legal y demás recaudos aplicables. No retirar los originales del Registro de Protocolizaciones para integrar esta carpeta.')
title(2,'2. Minuta de Automotores — esquema')
p('Ejemplo de datos concordantes. Deben trasladarse al formulario oficial vigente, completar sus restantes campos y suscribirlo según corresponda. Esta hoja no es una minuta oficial apta para ingreso.')
table([('Acto','Compraventa'),('Vendedora',v+'; cédula sintética '+vci+'; uruguaya; mayor de edad; soltera; domicilio ficticio Calle del Ejemplo 100, Montevideo'),('Comprador',c+'; cédula sintética '+cci+'; uruguayo; mayor de edad; soltero; domicilio ficticio Calle de la Demostración 200, Montevideo'),('Bien',vehicle),('Precio','USD 7.800'),('Documento','Contrato privado de 11/10/2026, certificado y protocolizado en el supuesto'),('Antecedente ilustrativo','Testimonio del título previo DEMO-ANT-001; inscripción DEMO-REG-001, inexistente'),('Protocolización ilustrativa','Número 12, folios 40 a 43 del año 2026, todos supuestos'),('Profesional',e),('Control SOA','Referencia POLIZA-DEMO-001; sin póliza real ni control acreditado')])
p('Debe corroborarse la concordancia de partes, estado civil completo, bien, antecedente, precio y referencias notariales entre formulario e instrumento. La ficha oficial de DGR contiene las instrucciones de presentación; el inicio digital no equivale a inscripción.')
title(3,'3. Documento privado de compraventa — ejemplo')
p(f'En Montevideo, el 11 de octubre de 2026, en el supuesto ficticio comparecen personalmente: por una parte, {v}, uruguaya, mayor de edad, soltera, cédula sintética {vci}, domiciliada ficticiamente en Calle del Ejemplo 100, Montevideo; y por otra parte, {c}, uruguayo, mayor de edad, soltero, cédula sintética {cci}, domiciliado ficticiamente en Calle de la Demostración 200, Montevideo. Se representa el siguiente contrato:')
p(f'PRIMERO. Objeto. {v}, parte vendedora, vende a {c}, parte compradora, quien compra, el siguiente bien: {vehicle}.')
p('SEGUNDO. Precio y pago. El precio supuesto es de siete mil ochocientos dólares estadounidenses (USD 7.800). A los solos efectos del ejemplo se representa una transferencia por ese importe, recibida por la vendedora, quien en el supuesto otorga carta de pago total. No existe comprobante ni recepción real de fondos; en una operación efectiva esta cláusula requiere evidencia de pago y revisión del régimen aplicable.')
p('TERCERO. Tradición. En el supuesto didáctico se representa la entrega del vehículo y sus llaves en esta fecha, a las 15:00 horas, con voluntad de transferir y adquirir propiedad y posesión. No se afirma entrega real; la cláusula definitiva debe corresponder al modo y hechos corroborados.')
p('CUARTO. Antecedente y situación del bien. Se supone para este ejemplo que la vendedora adquirió el bien por título previo DEMO-ANT-001, con inscripción DEMO-REG-001. Son referencias inexistentes. No se declara ausencia de gravámenes, afectaciones o deudas: su alcance exige certificados, antecedentes y controles reales.')
p('QUINTO. Controles. Situación fiscal y previsional, capacidad, titularidad, antecedentes, certificados, SOA y demás recaudos quedan sujetos a verificación profesional en una operación real. Este ejemplo no acredita esos extremos ni incluye declaraciones automáticas de no contribuyente.')
p(f'Suscripción simulada: {v} — SIN FIRMA REAL.\nSuscripción simulada: {c} — SIN FIRMA REAL.')
title(4,'4. Certificación notarial de firmas — ejemplo')
p(f'Fórmula simulada de otorgamiento y suscripción recibidos ante escribana, artículo 88 de la Ley 16.871. Personaje profesional: {e}. En el supuesto didáctico, las firmas del contrato precedente relativo al vehículo DEMO-001 serían puestas en su presencia por {v}, cédula sintética {vci}, y {c}, cédula sintética {cci}, quienes, previa lectura, lo ratificarían y suscribirían.')
p('No se certifica ningún hecho real. La identificación, capacidad y circunstancias de la comparecencia deben resultar de evidencia efectivamente examinada. No se presume conocimiento personal de los comparecientes. Si se ratifica un documento firmado con anterioridad, corresponde redactar esa actuación distinta.')
p('Procedencia ilustrativa: cláusula CUARTO del contrato del ejemplo. SOA: póliza ficticia POLIZA-DEMO-001 de ASEGURADORA EJEMPLO, cobertura y continuidad supuestas. Para inscripción real, controlar los artículos 27 y 28 de la Ley 18.412, incluido el período anterior exigible y el supuesto aplicable; esta referencia no acredita cobertura.')
p('Lugar y fecha ilustrativos: Montevideo, 11 de octubre de 2026. Destino supuesto: presentación registral de la compraventa. Papel, sello, signo y firma: no incorporados. Esta hoja no es una certificación autorizada.')
title(5,'5. Acta de protocolización preceptiva — ejemplo')
p(f'Número 12 — numeración ilustrativa. Montevideo, 11 de octubre de 2026. Se representa la siguiente agregación por {e} al Registro de Protocolizaciones, con fundamento en el artículo 89 vigente de la Ley 16.871:')
p('a) Contrato privado de compraventa de fecha 11 de octubre de 2026, relativo al vehículo ficticio DEMO-001, que en el supuesto ocupa dos fojas, folios 40 y 41; b) certificación notarial de sus firmas, de igual fecha, en una foja, folio 42; c) acta de agregación, en una foja, folio 43. Total supuesto: cuatro fojas. Estos folios no corresponden a una impresión real de este archivo.')
p('Referencia anterior ilustrativa: protocolización número 11, terminada en folio 39 del mismo año. No se supone que sea la primera protocolización. No se agrega poder porque el supuesto representa comparecencia personal; si existiera representación se deben verificar sus requisitos, incluido el artículo 291 de la Ley 18.362 cuando corresponda.')
p('Las fojas originales del supuesto quedarían en el Registro de Protocolizaciones. En una operación real deben verificarse documentos, orden, cantidad de fojas y carillas, numeración, referencias, actuación y formalidades antes de autorizar. Este texto no acredita incorporación efectiva. SIN SELLO, SIGNO NI FIRMA.')
title(6,'6. Testimonio notarial — estructura ilustrativa')
p('TESTIMONIO DE PROTOCOLIZACIÓN — ESQUEMA NO AUTORIZADO. La presentación registral de este supuesto se realizaría con el testimonio expedido en forma legal, acompañado por carátula, minuta y recaudos aplicables.')
p('Contenido que debe reproducirse con concordancia: 1) el documento privado de compraventa; 2) la certificación notarial de firmas; 3) el acta de protocolización. Las páginas 3, 4 y 5 del índice permiten ver esas piezas del supuesto; esta hoja indica la estructura y no contiene una reproducción literal integral autorizada.')
p(f'Cierre ilustrativo: se representaría la expedición por {e} de un testimonio concordante con la protocolización número 12, folios 40 a 43 del Registro de Protocolizaciones de 2026, destinado a la parte adquirente {c}, en Montevideo, en fecha corroborada. La escribana debe redactar la fórmula correspondiente y cumplir las formalidades de expedición. Aquí no hay firma, signo, sello, papel notarial validado ni actuación real.')
p('Antes de ventanilla: comprobar que el testimonio real sea completo y concordante; utilizar la minuta y carátula vigentes; cotejar antecedentes y demás recaudos; registrar luego la calificación y eventuales observaciones de DGR. Este ejemplo no debe presentarse como documento portante.')
p('Fuentes y alcance: Ley 16.871, artículos 25, 88 y 89; Ley 18.412, artículos 27 y 28; ficha oficial de inscripción de automotores en gub.uy. Manual de Automotores 2025 de DGR localizado, con contenido pendiente de consulta. Ver informe legal del proyecto.')
for attr in ('author','last_modified_by','comments','keywords'):setattr(d.core_properties,attr,'')
path=ROOT/'ejemplos/Carpeta_CV_automotor_EJEMPLO_FICTICIO.docx';d.save(path)
print('Ejemplo ficticio generado:',path)
