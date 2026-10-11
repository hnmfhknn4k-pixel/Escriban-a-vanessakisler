"""Carga local y controles con documentos ficticios; sin subidas al servidor."""
import io,zipfile,hashlib
from docx import Document
from playwright.sync_api import sync_playwright,expect

def docx():
 d=Document();d.add_paragraph('Comparece JUAN FICTICIO, cédula 1.234.567-3. [PRECIO_PENDIENTE] la la')
 d.add_table(rows=1,cols=1).cell(0,0).text='TABLA FICTICIA'
 d.sections[0].header.paragraphs[0].text='ENCABEZADO FICTICIO'
 d.sections[0].footer.paragraphs[0].text='PIE FICTICIO'
 out=io.BytesIO();d.save(out);return out.getvalue()
def altered(blob,name,data,duplicate=False):
 out=io.BytesIO()
 with zipfile.ZipFile(io.BytesIO(blob)) as source,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  for item in source.infolist():
   if item.filename!=name or duplicate:z.writestr(item,source.read(item.filename))
  z.writestr(name,data)
 return out.getvalue()
with sync_playwright() as p:
 browser=p.chromium.launch(executable_path='/usr/bin/chromium',headless=True,args=['--no-sandbox'])
 page=browser.new_page(viewport={'width':390,'height':844});errors=[];requests=[]
 page.on('pageerror',lambda e:errors.append(str(e)));page.on('request',lambda r:requests.append((r.method,r.url)))
 page.goto('http://127.0.0.1:8001/demo');input=page.locator('#review-file')
 def load(name,blob,mime='application/octet-stream'):input.set_input_files({'name':name,'mimeType':mime,'buffer':blob})
 text='Texto ficticio: 1.234.567-2, 1.234.567-3. [PENDIENTE] la la <script>alert(1)</script>'.encode()
 load('ficticio.txt',text,'text/plain');expect(page.locator('#file-state')).to_contain_text('Archivo leído localmente')
 expect(page.locator('#file-hashes')).to_contain_text(hashlib.sha256(text).hexdigest());page.get_by_role('button',name='Validar archivo cargado').click()
 expect(page.locator('#file-checks')).to_contain_text('no acredita identidad');expect(page.locator('#file-checks')).to_contain_text('Dígito verificador incompatible')
 expect(page.locator('#file-checks')).to_contain_text('Posible repetición');expect(page.locator('#file-text')).to_contain_text('<script>alert(1)</script>')
 with page.expect_download() as download:page.get_by_role('button',name='Descargar revisión TXT').click()
 report=__import__('pathlib').Path(download.value.path()).read_text();assert 'Registro editable' in report;assert hashlib.sha256(text).hexdigest() in report
 blob=docx();load('ficticio.docx',blob);expect(page.locator('#file-state')).to_contain_text('Archivo leído localmente')
 for s in ('TABLA FICTICIA','ENCABEZADO FICTICIO','PIE FICTICIO'):expect(page.locator('#file-text')).to_contain_text(s)
 expect(page.locator('#file-hashes')).to_contain_text(hashlib.sha256(blob).hexdigest());expect(page.locator('#file-sections')).to_contain_text('No se verifica adecuación normativa')
 assert page.get_by_role('button',name='Descargar revisión TXT').is_disabled()
 page.get_by_role('button',name='Validar archivo cargado').click();expect(page.locator('#file-checks')).to_contain_text('Campo pendiente')
 load('duplicado.docx',altered(blob,'word/document.xml',b'<broken/>',duplicate=True));expect(page.locator('#file-state')).to_contain_text('duplicadas');assert page.locator('#file-result').is_hidden()
 load('entidades.docx',altered(blob,'word/document.xml',b'<!DOCTYPE x><x/>'));expect(page.locator('#file-state')).to_contain_text('entidades no admitidas')
 load('bomba.docx',altered(blob,'oversized.bin',b'0'*(41*1024*1024)));expect(page.locator('#file-state')).to_contain_text('excede el límite')
 load('corrupto.docx',b'not-a-zip');expect(page.locator('#file-state')).to_contain_text('DOCX inválido')
 load('antiguo.doc',b'not-word');expect(page.locator('#file-state')).to_contain_text('Admitimos DOCX y TXT')
 load('grande.txt',b'a'*(8*1024*1024+1));expect(page.locator('#file-state')).to_contain_text('hasta 8 MB')
 load('ultimo.txt',b'Ficticio [CAMPO]');expect(page.locator('#file-state')).to_contain_text('Archivo leído localmente')
 page.get_by_role('button',name='Quitar archivo').click();expect(page.locator('#file-text')).to_have_text('');assert page.locator('#file-result').is_hidden()
 assert not any(method!='GET' or '/api/' in url for method,url in requests)
 assert page.evaluate('() => document.documentElement.scrollWidth <= innerWidth');assert not errors,errors
 browser.close();print('Carga local DOCX/TXT, tablas/encabezados/pies, huellas, controles, descarga, ZIP inválido/duplicado/excesivo, entidades y retiro: OK')
