import base64
import io
import json
import os
import secrets
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from docx import Document
from fastapi.testclient import TestClient
import httpx

TEMP=tempfile.TemporaryDirectory();os.environ['ESCRIBANIA_DATA']=TEMP.name
import app as service
from core import *
from store import Store,password_hash,totp
from ai import assist,readiness
from profile import validate_profile

PARTY={'name':'MARÍA MUÑOZ PÉREZ','ci':'1.234.567-2'}
EVIDENCE={'reason':'Cotejo ficticio documentado para prueba','evidence':'Documento ficticio de laboratorio'}
SECRET='JBSWY3DPEHPK3PXP'

def document(text='Comparece MARÍA MUÑOZ PÉREZ, cédula 1.234.567-2.'):
    d=Document();d.add_paragraph(text);out=io.BytesIO();d.save(out);return out.getvalue()

def profile():
    # Synthetic test measures, never installed as normative defaults.
    return {'w':210,'h':297,'top':30,'bottom':30,'left':30,'right':30,'header':10,'footer':10,'gutter':0,'source':'https://www.impo.com.uy/test-only','authority':'Laboratorio ficticio','normative_version':'SOLO PRUEBA FICTICIA','evidence':'Prueba ficticia no normativa','reviewer':'Revisor ficticio','validated_at':'2026-10-11','print_test':'Laboratorio ficticio, anverso y reverso','duplex_validated':True,'paper':'liso','mirror':True}

class CoreTests(unittest.TestCase):
    def test_uppercase_accents_and_n(self): self.assertEqual(upper_name('  María Muñoz Pérez '),'MARÍA MUÑOZ PÉREZ')
    def test_unicode_composition(self): self.assertEqual(upper_name('Pen\u0303a'),'PEÑA')
    def test_valid_ci_formats(self):
        for s in ('1.234.567-2','1234567-2','12345672'): self.assertTrue(valid_ci(s))
    def test_ci_wrong_check_or_format(self):
        for s in ('1.234.567-3','123','123A45672','1.23.4567-2'):self.assertFalse(valid_ci(s))
    def test_reversible_redaction(self):
        text=PARTY['name']+' '+PARTY['ci']
        hidden,mapping,residual=pseudonymize(text,[PARTY])
        self.assertNotIn('MARÍA',hidden);self.assertNotIn('1.234.567',hidden)
        self.assertEqual(reinsert(hidden,mapping),text);self.assertFalse(residual)
    def test_alternative_ci_format_hidden(self):
        hidden,m,_=pseudonymize('MARÍA MUÑOZ PÉREZ 12345672',[PARTY]);self.assertNotIn('12345672',hidden)
        self.assertIn(PARTY['ci'],reinsert(hidden,m))
    def test_unknown_reference_rejected(self):
        with self.assertRaises(ValueError):reinsert('[[PERSONA_9_NAME]]',{})
    def test_reserved_delimiter_rejected(self):
        with self.assertRaises(ValueError):pseudonymize('[[INSTRUCCION]]',[PARTY])
    def test_indirect_identifiers_detected(self):
        hidden,m,r=pseudonymize('Correo ana@example.test; domicilio Calle Ficticia 123',[PARTY]);self.assertIn('correo',r);self.assertIn('ubicación u otro identificador',r)
    def test_missing_and_wrong_mentions(self):
        self.assertTrue(identity_issues('MARÍA MUÑOZ PÉREZ, cédula 1.234.567-3',[PARTY]))
        self.assertTrue(identity_issues('María Muñoz Pérez, cédula 1.234.567-2',[PARTY]))
        self.assertFalse(identity_issues('Comparece MARÍA MUÑOZ PÉREZ, cédula 1.234.567-2.',[PARTY]))
    def test_unknown_person_detection(self):
        self.assertTrue(any('no cotejada' in x for x in identity_issues('Comparece JUAN FICTICIO, cédula 1.234.567-2.',[PARTY])))
    def test_split_runs(self):
        d=Document();p=d.add_paragraph();p.add_run('MARÍA ');p.add_run('MUÑOZ PÉREZ, 1.234.567-2');out=io.BytesIO();d.save(out)
        segments,w=read_docx(out.getvalue());self.assertFalse(identity_issues(analyzed_text(segments),[PARTY]))
    def test_tables_headers_footers_preserved(self):
        d=Document();d.add_paragraph('Principal');d.add_table(1,1).cell(0,0).text='Tabla';d.sections[0].header.paragraphs[0].text='Encabezado';d.sections[0].footer.paragraphs[0].text='Pie';out=io.BytesIO();d.save(out)
        segments,w=read_docx(out.getvalue());text=analyzed_text(segments)
        for t in ('Principal','Tabla','Encabezado','Pie'):self.assertIn(t,text)
        final=export_docx(out.getvalue(),[],'revisor')
        self.assertEqual(analyzed_text(read_docx(final)[0]),text)
        with zipfile.ZipFile(io.BytesIO(final)) as z:self.assertIn('word/styles.xml',z.namelist())
    def test_tracked_export(self):
        original=document('texto original');seg=read_docx(original)[0][0]
        final=export_docx(original,[{'part':seg['part'],'index':seg['index'],'before':'texto original','after':'texto revisado','status':'accepted'}],'revisor')
        with zipfile.ZipFile(io.BytesIO(final)) as z:
            xml=z.read('word/document.xml').decode();self.assertIn('w:ins',xml);self.assertIn('w:del',xml);self.assertIn('texto original',xml);self.assertIn('texto revisado',xml)
        self.assertIn('texto revisado',analyzed_text(read_docx(final)[0]))
        self.assertTrue(read_docx(final)[1])
    def test_invalid_zip(self):
        with self.assertRaises(ValueError):read_docx(b'bad')
    def test_oversized(self):
        with self.assertRaises(ValueError):read_docx(b'x'*(10*1024*1024+1))
    def test_profile_requires_source_and_evidence(self):
        with self.assertRaises(ValueError):validate_profile({'w':210})
    def test_profile_rejects_nan(self):
        p=profile();p['top']=float('nan')
        with self.assertRaises(ValueError):validate_profile(p)

class AuditTests(unittest.TestCase):
    def setUp(self):self.tmp=tempfile.TemporaryDirectory();self.s=Store(self.tmp.name);self.c={'id':'ficticio','name':'MARÍA MUÑOZ'};self.s.save(self.c,'revisor','create',{},True)
    def tearDown(self):self.tmp.cleanup()
    def test_audit_identity(self):self.assertEqual(self.s.history('ficticio','revisor')[0]['actor'],'revisor')
    def test_no_plaintext_on_disk(self):self.assertFalse('MARÍA MUÑOZ'.encode() in self.s.db.read_bytes())
    def test_update_and_delete_blocked(self):
        import sqlite3
        with self.s.connect() as db:
            with self.assertRaises(sqlite3.IntegrityError):db.execute('UPDATE audit SET mac=?',('fake',))
            with self.assertRaises(sqlite3.IntegrityError):db.execute('DELETE FROM audit')
    def test_tampered_mac_detected(self):
        with self.s.connect() as db:db.execute('DROP TRIGGER no_audit_update');db.execute("UPDATE audit SET mac='fake'")
        with self.assertRaises(ValueError):self.s.verify()
    def test_payload_rollback_detected(self):
        with self.s.connect() as db:old=db.execute('SELECT payload FROM cases').fetchone()[0]
        self.c['name']='Nuevo';self.s.save(self.c,'revisor','changed',{})
        with self.s.connect() as db:db.execute('UPDATE cases SET payload=?',(old,))
        with self.assertRaises(ValueError):self.s.get('ficticio','revisor')
    def test_case_permissions(self):
        with self.assertRaises(PermissionError):self.s.get('ficticio','otro')

class ApiTests(unittest.TestCase):
    def setUp(self):
        self.client=TestClient(service.app);self.name='u'+secrets.token_hex(6)
        with service.store.connect() as db:db.execute('INSERT INTO users VALUES(?,?,?)',(self.name,password_hash('ficticio-123456789'),service.store.fernet.encrypt(SECRET.encode())))
        self.origin={'Origin':service.ORIGIN};r=self.client.post('/api/login',json={'user':self.name,'password':'ficticio-123456789','otp':totp(SECRET,time.time())},headers=self.origin)
        self.assertEqual(r.status_code,200,r.text);self.headers=self.origin|{'X-CSRF-Token':r.json()['csrf']}
    def post(self,path,data):return self.client.post(path,json=data,headers=self.headers)
    def import_case(self,text=None):
        r=self.client.post('/api/cases',files={'file':('ficticio.docx',document(text) if text else document())},headers=self.headers);self.assertEqual(r.status_code,200,r.text);return r.json()
    def party(self,c):return self.post('/api/cases/'+c['id']+'/parties',PARTY|EVIDENCE)
    def test_login_mfa_required(self):
        r=self.client.post('/api/login',json={'user':self.name,'password':'ficticio-123456789','otp':'000000'},headers=self.origin);self.assertEqual(r.status_code,401)
    def test_csrf_required(self):self.assertEqual(self.client.post('/api/draft',json={'title':'Ficticio'},headers=self.origin).status_code,403)
    def test_other_origin_denied(self):self.assertEqual(self.client.post('/api/draft',json={'title':'Ficticio'},headers=self.headers|{'Origin':'https://evil.test'}).status_code,403)
    def test_import_party_analysis(self):
        c=self.import_case();r=self.party(c);self.assertEqual(r.status_code,200,r.text)
        r=self.post('/api/cases/'+c['id']+'/analyze',{});self.assertEqual(r.status_code,200);self.assertEqual(r.json()['parties'][0]['name'],PARTY['name']);self.assertIn('engine',r.json()['analysis'])
    def test_identity_change_blocked(self):
        c=self.import_case();self.party(c);s=c['segments'][0]
        r=self.post('/api/cases/'+c['id']+'/changes',EVIDENCE|{'segment':s['id'],'after':s['text'].replace('MUÑOZ','PÉREZ')});self.assertEqual(r.status_code,409)
    def test_ci_change_blocked(self):
        c=self.import_case();self.party(c);s=c['segments'][0]
        r=self.post('/api/cases/'+c['id']+'/changes',EVIDENCE|{'segment':s['id'],'after':s['text'].replace('1.234.567-2','1.234.567-3')});self.assertEqual(r.status_code,409)
    def test_changes_accept_reject_undo(self):
        c=self.import_case();self.party(c);s=c['segments'][0];path='/api/cases/'+c['id']
        r=self.post(path+'/changes',EVIDENCE|{'segment':s['id'],'after':s['text']+' Texto ficticio.'});self.assertEqual(r.status_code,200,r.text);ch=r.json()['changes'][0]
        r=self.post(path+'/changes/'+ch['id'],EVIDENCE|{'status':'accepted'});self.assertIn('Texto ficticio',r.json()['working'][0]['text'])
        r=self.post(path+'/changes/'+ch['id'],EVIDENCE|{'status':'undone'});self.assertEqual(r.json()['working'][0]['text'],s['text'])
        r=self.post(path+'/changes',EVIDENCE|{'segment':s['id'],'after':s['text']+' Otra propuesta.'});ch=r.json()['changes'][-1]
        r=self.post(path+'/changes/'+ch['id'],EVIDENCE|{'status':'rejected'});self.assertEqual(r.json()['working'][0]['text'],s['text'])
    def test_final_approval_blocked_without_template(self):
        c=self.import_case();self.party(c);r=self.post('/api/cases/'+c['id']+'/approve',EVIDENCE);self.assertEqual(r.status_code,409);self.assertIn('Perfil normativo',r.text)
    def test_download_original_exact(self):
        blob=document();r=self.client.post('/api/cases',files={'file':('ficticio.docx',blob)},headers=self.headers);cid=r.json()['id']
        r=self.client.get('/api/cases/'+cid+'/download/original');self.assertEqual(r.content,blob)
    def test_psuedonym_preview_no_transfer(self):
        c=self.import_case();self.party(c);r=self.post('/api/cases/'+c['id']+'/pseudonymize',{'extra_terms':[]});self.assertFalse(r.json()['outbound']);self.assertNotIn(PARTY['name'],r.json()['text'])
    def test_model_unconfigured(self):
        c=self.import_case();r=self.post('/api/cases/'+c['id']+'/assist',{});self.assertEqual(r.status_code,503)
    def test_split_name_cannot_be_changed(self):
        d=Document();p=d.add_paragraph();p.add_run('MARÍA ');p.add_run('MUÑOZ PÉREZ, 1.234.567-2');out=io.BytesIO();d.save(out)
        r=self.client.post('/api/cases',files={'file':('split-ficticio.docx',out.getvalue())},headers=self.headers);c=r.json();self.party(c)
        s=c['segments'][1];r=self.post('/api/cases/'+c['id']+'/changes',EVIDENCE|{'segment':s['id'],'after':'PÉREZ PÉREZ, 1.234.567-2'});self.assertEqual(r.status_code,409)
    def test_identity_discrepancy_corrected_only_from_ficha(self):
        c=self.import_case('Comparece MARÍA MUÑOZ PÉREZ, cédula 1.234.567-3.');c=self.party(c).json();path='/api/cases/'+c['id']
        r=self.post(path+'/identity-fix',EVIDENCE|{'party':c['parties'][0]['id'],'field':'ci','segment':c['segments'][0]['id'],'old':'1.234.567-3'})
        self.assertEqual(r.status_code,200,r.text);change=r.json()['changes'][0];self.assertEqual(change['status'],'pending');self.assertIn('1.234.567-2',change['after'])
        c=self.post(path+'/changes/'+change['id'],EVIDENCE|{'status':'accepted'}).json();self.assertFalse(identity_issues(analyzed_text(c['working']),c['parties']))
    def test_confirmed_identity_cannot_be_reassigned(self):
        c=self.import_case();c=self.party(c).json();r=self.post('/api/cases/'+c['id']+'/identity-fix',EVIDENCE|{'party':c['parties'][0]['id'],'field':'name','segment':c['segments'][0]['id'],'old':'MUÑOZ'});self.assertEqual(r.status_code,409)
    def test_pending_fields_stay_blocked(self):
        c=self.import_case('Comparece MARÍA MUÑOZ PÉREZ, cédula 1.234.567-2. Precio [PRECIO].');self.assertTrue(any('corchetes' in b for b in c['blockers']))
    def test_print_review_required_when_profile_installed(self):
        with patch.object(service,'GLOBAL_PROFILE',profile()):
            c=self.import_case();self.party(c);path='/api/cases/'+c['id'];self.post(path+'/analyze',{})
            r=self.post(path+'/corroborate',EVIDENCE|{'all_mentions_checked':True});self.assertEqual(r.status_code,422)
    def test_global_profile_cannot_be_changed_per_case(self):
        c=self.import_case();self.assertEqual(self.post('/api/cases/'+c['id']+'/template',profile()).status_code,404)
    def test_final_and_reopening_fictitious_profile(self):
        with patch.object(service,'GLOBAL_PROFILE',profile()):
            c=self.import_case();self.party(c);path='/api/cases/'+c['id'];r=self.post(path+'/analyze',{});c=r.json()
            for o in c['observations']:self.post(path+'/observations/'+o['id'],EVIDENCE|{'decision':'rejected'})
            self.post(path+'/corroborate',EVIDENCE|{'all_mentions_checked':True,'print_checked':True,'reviewed_sha256':c['reviewed_hash']})
            r=self.post(path+'/approve',EVIDENCE);self.assertEqual(r.status_code,200,r.text);c=r.json();self.assertTrue(c['approved'])
            r=self.client.get(path+'/download/final');self.assertEqual(digest(r.content),c['final_hash'])
            r=self.post(path+'/reopen',EVIDENCE);self.assertFalse(r.json()['approved']);self.assertTrue(r.json()['blockers'])

class AiTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.text='Comparece MARÍA MUÑOZ PÉREZ, cédula 1.234.567-2.'
        review={f:'Condición revisada por responsable' for f in ('contract','training','retention','subprocessors','international_transfers','legal_basis','reviewer','reviewed_at')};review.update(approved=True,expires_at='2099-01-01T00:00:00+00:00');Path(self.temp.name,'provider-review.json').write_text(json.dumps(review))
    def tearDown(self):self.temp.cleanup()
    def result(self,kind='inferencia',fragment='Comparece'):
        return {'observations':[dict(kind=kind,finding='Revisar el acto',fragment=fragment,basis='Falta evidencia del acto',alternatives='Corroborar documentación',uncertainty='No se determina',reconsider='Aportar documentos')],'draft':''}
    def run_model(self,result):
        def handler(request):
            body=json.loads(request.content);self.assertNotIn(PARTY['name'],json.dumps(body));self.assertNotIn(PARTY['ci'],json.dumps(body))
            return httpx.Response(200,json={'stop_reason':'end_turn','model':'modelo-ficticio','content':[{'type':'text','text':json.dumps(result)}]})
        hidden,_,_=pseudonymize(self.text,[PARTY])
        with patch.dict(os.environ,{'ANTHROPIC_API_KEY':'clave-ficticia','ANTHROPIC_MODEL':'modelo-ficticio'}):return assist(self.temp.name,self.text,[PARTY],[],'revisor',digest(hidden),True,transport=httpx.MockTransport(handler))
    def test_outbound_contains_references_only(self):self.assertEqual(self.run_model(self.result())['metadata']['prompt_version'],'notarial-0.4-private-1')
    def test_model_cannot_claim_facts(self):
        with self.assertRaises(ValueError):self.run_model(self.result('hecho'))
    def test_invented_quote_rejected(self):
        with self.assertRaises(ValueError):self.run_model(self.result(fragment='Texto que no existe'))
    def test_unknown_reference_rejected(self):
        result=self.result();result['draft']='[[PERSONA_99_NAME]]'
        with self.assertRaises(ValueError):self.run_model(result)
    def test_unverified_law_rejected(self):
        result=self.result();result['observations'][0]['finding']='Ley 99999 artículo 99'
        with self.assertRaises(ValueError):self.run_model(result)

if __name__=='__main__':unittest.main()
