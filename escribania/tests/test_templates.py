import json
import re
import unittest
from pathlib import Path
from docx import Document
from core import read_docx,analyzed_text,digest,NS,W
from templates import load_template,catalog
import test_app as baseline

class TemplateDocumentTests(unittest.TestCase):
    def test_template_derived_from_supplied_model(self):
        manifest,blob=load_template('compraventa-automotor')
        self.assertEqual(digest(blob),manifest['template_sha256'])
        self.assertTrue(manifest['review_required'])
        self.assertFalse(manifest['normative_format_verified'])
        self.assertEqual(len(manifest['changes']),14)
        self.assertEqual(manifest['source_filename'],'MODELO CV AUTOMOTOR.doc')
    def test_placeholders_replace_assumed_case_data(self):
        manifest,blob=load_template('compraventa-automotor');text=analyzed_text(read_docx(blob)[0])
        for value in ('19 de setiembre','2025','siete mil ochocientos','16.30','5190'):
            self.assertNotIn(value,text)
        self.assertIsNone(re.search(r'x{3,}',text,re.I))
        self.assertIn('[VENDEDOR_NOMBRES_APELLIDOS]',text)
        self.assertIn('[COMPRADOR_NOMBRES_APELLIDOS]',text)
        self.assertIn('[VEHICULO_CHASIS]',text)
        self.assertIn('[FORMA_Y_EVIDENCIA_DE_PAGO]',text)
    def test_metadata_and_false_tab_warning(self):
        manifest,blob=load_template('compraventa-automotor');doc=Document(__import__('io').BytesIO(blob))
        self.assertFalse(doc.core_properties.author);self.assertFalse(doc.core_properties.last_modified_by)
        self.assertFalse(read_docx(blob)[1])
    def test_actual_tab_still_warns(self):
        import io
        doc=Document();doc.add_paragraph('Texto\tcon tabulación');out=io.BytesIO();doc.save(out)
        self.assertTrue(any('tab' in w for w in read_docx(out.getvalue())[1]))
    def test_adaptations_match_prepared_document(self):
        manifest,blob=load_template('compraventa-automotor');doc=Document(__import__('io').BytesIO(blob))
        for change in manifest['changes']:self.assertEqual(doc.paragraphs[change['paragraph']].text,change['after'])

class TemplateRoutesTests(unittest.TestCase):
    setUp=baseline.ApiTests.setUp
    post=baseline.ApiTests.post
    def test_catalog_authenticated(self):
        response=self.client.get('/api/templates');self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()[0]['id'],'compraventa-automotor')
        self.assertGreater(len(response.json()[0]['fields']),30)
    def test_acknowledgment_required(self):
        response=self.post('/api/templates/compraventa-automotor/cases',{})
        self.assertEqual(response.status_code,422)
    def test_creation_remains_unapproved_and_has_provenance(self):
        response=self.post('/api/templates/compraventa-automotor/cases',{'review_required_acknowledged':True})
        self.assertEqual(response.status_code,200,response.text);case=response.json()
        self.assertFalse(case['approved']);self.assertFalse(case['parties'])
        self.assertTrue(any('corchetes' in b for b in case['blockers']))
        self.assertEqual(case['original_hash'],case['template_origin']['template_sha256'])
        response=self.post('/api/cases/'+case['id']+'/approve',baseline.EVIDENCE);self.assertEqual(response.status_code,409)
        history=self.client.get('/api/cases/'+case['id']+'/audit').json()['events']
        self.assertEqual(history[0]['action'],'template_case_created')
    def test_download_exact_template(self):
        _,blob=load_template('compraventa-automotor')
        response=self.client.get('/api/templates/compraventa-automotor/download')
        self.assertEqual(response.content,blob)
    def test_unauthenticated_catalog_and_unknown_template(self):
        from fastapi.testclient import TestClient
        import app
        with TestClient(app.app) as anonymous:self.assertEqual(anonymous.get('/api/templates').status_code,401)
        self.assertEqual(self.client.get('/api/templates/no-existe/download').status_code,404)

if __name__=='__main__':unittest.main()
